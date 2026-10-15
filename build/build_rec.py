"""Records / Excel-export columns, laid out like a reference Mapper export (column set and order only).

The reference file (e.g. PWEIII - Mapper Export - 03.24.26.xlsx) supplies ONLY the header row.
Values come from the project's current export, aligned row-for-row with the dashboard data
(rows added by hand to the dashboard data get their values from that data).
usage: python3 build_rec.py <reference.xlsx> <current_export.xlsx> <data.json> <out.json>
"""
import datetime, json, sys
import openpyxl

ref, cur, data_path, out = sys.argv[1:5]
G = [c.value for c in next(openpyxl.load_workbook(ref, read_only=True).worksheets[0].iter_rows(max_row=1))]
ws = openpyxl.load_workbook(cur, read_only=True, data_only=True).worksheets[0]
rows = list(ws.iter_rows(values_only=True)); H = rows[0]
E = [dict(zip(H, r)) for r in rows[1:] if any(r)]
cols = [g for g in G if g in H]                      # reference order, only columns this export has
cols = ['BRK_LSE_NO'] + [c for c in cols if c != 'BRK_LSE_NO']   # contract no. first (user request)
D = json.load(open(data_path))

def val(v):
    if isinstance(v, (datetime.date, datetime.datetime)): return v.strftime('%Y-%m-%d')
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15: return int(v)
    return v
# hand-added rows (no export row): fill what the dashboard row knows
FROM_DATA = {'TR_NUM': 'id', 'CLIENT_UNIT': 'unit', 'ABSTRACT': 'abs', 'TRACT_COUNTY': 'cty', 'TRACT_ACRES': 'tac', 'MIN_INT': 'int',
             'NET_AC': 'nac', 'ASSOCIATEID': 'oid', 'MINERAL_OWNER': 'owner', 'ADDRESS': 'addr', 'PHONE_NO': 'phone', 'HELD_BY': 'held',
             'LEASING_STATUS': 'ls', 'LESSOR': 'lessor', 'CURRENT_LESSEE': 'lessee', 'BRK_LSE_NO': 'lno', 'SYMBOL': 'sym', 'VOL_PAGE': 'vp',
             'EFFECTIVE_DATE': 'eff', 'EXPIRATION_DATE': 'exp', 'LSE_GROSS_AC': 'lgac', 'LSE_NET_AC': 'lnac', 'BONUS_PER_AC': 'bonus',
             'ROYALTY': 'roy', 'TERM': 'term'}
U = {}; uv = []
def code(v):
    k = json.dumps(v)
    if k not in U: U[k] = len(uv); uv.append(v)
    return U[k]
V = []; j = 0; manual = 0
for r in D:
    e = E[j] if j < len(E) else None
    if e is not None and e['TR_NUM'] == r['id'] and e['ASSOCIATEID'] == r['oid']:
        V.append([code(val(e.get(c))) for c in cols]); j += 1
    else:
        manual += 1
        V.append([code(r.get(FROM_DATA[c]) if c in FROM_DATA else None) for c in cols])
assert j == len(E), f'only {j} of {len(E)} export rows matched'
json.dump({'cols': cols, 'u': uv, 'v': V}, open(out, 'w'), separators=(',', ':'))
print(out, len(cols), 'columns |', len(V), 'rows (', manual, 'added by hand ) |', len(uv), 'distinct values | missing from export:', [g for g in G if g not in H])
